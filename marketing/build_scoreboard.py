"""
Builds marketing/TrueNorth-Growth-Scoreboard.xlsx — door economics, a weekly
scoreboard, and cost-per-door by channel. All results are live formulas.

    <python with openpyxl> marketing/build_scoreboard.py
"""
import datetime as dt
import pathlib

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

OUT = pathlib.Path(__file__).resolve().parent / "TrueNorth-Growth-Scoreboard.xlsx"

NAVY, GOLD, SAND, GREY = "0F2742", "C29A4A", "F6F1E7", "808080"
F = "Arial"
title_font = Font(name=F, size=16, bold=True, color=NAVY)
sub_font = Font(name=F, size=10, italic=True, color="52627A")
head_font = Font(name=F, size=10, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor=NAVY)
section_font = Font(name=F, size=11, bold=True, color=NAVY)
section_fill = PatternFill("solid", fgColor=SAND)
input_font = Font(name=F, size=10, color="0000FF")  # blue = your inputs
calc_font = Font(name=F, size=10, color="000000")  # black = formulas
label_font = Font(name=F, size=10, color="000000")
note_font = Font(name=F, size=9, italic=True, color="52627A")
example_font = Font(name=F, size=10, italic=True, color=GREY)
key_fill = PatternFill("solid", fgColor="FFFF00")  # yellow = key assumption to confirm
input_fill = PatternFill("solid", fgColor="FFF9E6")
thin = Side(style="thin", color="D9D2C3")
box = Border(bottom=thin)

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'

wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)


def header_row(ws, row, labels, widths=None):
    for i, text in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=text)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center" if i > 1 else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def section(ws, row, text, span=3):
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = section_fill
    c = ws.cell(row=row, column=1, value=text)
    c.font = section_font


# ── Read me ──────────────────────────────────────────────────────────────────
rm = wb.active
rm.title = "Read me"
rm.column_dimensions["A"].width = 110
lines = [
    ("TrueNorth Growth Scoreboard", title_font),
    ("One number decides every marketing choice: what it costs to sign one door (property), compared with what a door is worth.", sub_font),
    ("", None),
    ("How to read the colors", section_font),
    ("Blue text = an input you type or change.   Black text = a formula (don't overwrite).   Yellow fill = a key assumption to confirm with your own numbers.", label_font),
    ("", None),
    ("Tabs", section_font),
    ("Door Economics — what one door earns per year and over its life, the most you can spend to win one, doors needed for your income goal, and a postcard calculator.", label_font),
    ("Weekly Scoreboard — fill one row every Monday (about 5 minutes). The grey 'Example' row shows the format and is excluded from totals.", label_font),
    ("Channels — spend, leads, and signed doors per channel. Cost per door is compared against the maximum from Door Economics.", label_font),
    ("", None),
    ("Where the numbers come from", section_font),
    ("Leads: every website lead email now ends with its channel in [brackets], e.g. [postcard / direct_mail], [google_business_profile / organic], [direct].", label_font),
    ("Spend: postcard invoices, ad platform billing, printing. Reviews: your Google Business Profile. Doors: signed management agreements.", label_font),
    ("", None),
    ("Starting values are placeholders from the website's 'starting at' pricing and common planning assumptions — replace them with your real figures.", note_font),
]
for i, (text, font) in enumerate(lines, 1):
    c = rm.cell(row=i, column=1, value=text)
    if font:
        c.font = font
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ── Door Economics ───────────────────────────────────────────────────────────
de = wb.create_sheet("Door Economics")
de.column_dimensions["A"].width = 52
de.column_dimensions["B"].width = 16
de.column_dimensions["C"].width = 78
de["A1"], de["A1"].font = "Door Economics", title_font
de["A2"], de["A2"].font = "What one managed property (door) is worth, and what you can afford to spend to win one.", sub_font

rows = [
    # (row, label, value_or_formula, fmt, is_input, key_assumption, note)
    (4, "Inputs", None, None, None, None, None),
    (5, "Average monthly management fee ($)", 149, USD, True, True, "Plus plan 'starting at' price on the website (src/lib/pricing.ts). Use your real average."),
    (6, "Average monthly rent ($)", 2200, USD, True, True, "Example Spring Hill single-family rent. Replace with your portfolio average."),
    (7, "Leasing fee (share of one month's rent)", 0.75, PCT, True, True, "Leasing 'starting at' 75% of first month (website). Confirm Exhibit B line 3."),
    (8, "Tenant turnover rate (per year)", 0.5, PCT, True, True, "Assumption: 50% = a new tenant about every 2 years."),
    (9, "Renewal fee ($)", 299, USD, True, False, "Renewal 'starting at' $299 (website). Charged in years without turnover."),
    (10, "Direct cost per door ($ per year)", 300, USD, True, True, "Assumption: software, mileage, supplies per door. Replace with actuals."),
    (11, "Average years an owner stays", 3.5, NUM1, True, True, "Assumption. Retention is the biggest lever on lifetime value."),
    (12, "Max share of lifetime contribution to spend winning a door", 0.2, PCT, True, False, "Assumption: keep acquisition cost at or under 20% of lifetime contribution."),
    (13, "Owner income goal ($ per year)", 100000, USD, True, False, "Your target take-home before tax."),
    (14, "Fixed overhead ($ per year)", 12000, USD, True, True, "Assumption: E&O and liability insurance, licenses, software base fees, baseline marketing."),
    (16, "Per-door results", None, None, None, None, None),
    (17, "Management fees per door per year", "=B5*12", USD, False, False, None),
    (18, "Leasing fees per door per year", "=B6*B7*B8", USD, False, False, "Rent x leasing % x how often a new tenant is placed."),
    (19, "Renewal fees per door per year", "=B9*(1-B8)", USD, False, False, "Renewal fee in the years the tenant stays."),
    (20, "Revenue per door per year", "=SUM(B17:B19)", USD, False, False, None),
    (21, "Contribution per door per year", "=B20-B10", USD, False, False, "Revenue minus direct cost."),
    (22, "Lifetime contribution per door", "=B21*B11", USD, False, False, None),
    (23, "Max you can spend to win one door", "=B22*B12", USD, False, False, "Compare every channel's cost per door to this (Channels tab)."),
    (25, "Scale", None, None, None, None, None),
    (26, "Doors needed for your income goal", '=IF(B21>0,ROUNDUP((B13+B14)/B21,0),"-")', NUM, False, False, "(Income goal + overhead) / contribution per door."),
    (27, "Annual revenue at that door count", '=IF(ISNUMBER(B26),B26*B20,"-")', USD, False, False, None),
    (29, "Postcard calculator", None, None, None, None, None),
    (30, "Postcards mailed", 1000, NUM, True, False, "One drop to the absentee-owner list."),
    (31, "All-in cost per postcard (print + postage + list)", 0.75, USD2, True, True, "Assumption: 6x9 marketing mail commonly lands around $0.60–$1.00 all-in. Get a vendor quote."),
    (32, "Response rate (leads per postcard)", 0.01, PCT, True, True, "Assumption: 0.5%–1.5% is a common planning range; repeat drops usually respond better."),
    (33, "Lead-to-signed-door rate", 0.2, PCT, True, True, "Assumption: 1 in 5 owners who request an analysis signs."),
    (34, "Total mail cost", "=B30*B31", USD, False, False, None),
    (35, "Leads expected", "=B30*B32", NUM1, False, False, None),
    (36, "Doors expected", "=B35*B33", NUM1, False, False, None),
    (37, "Cost per door from postcards", '=IF(B36>0,B34/B36,"-")', USD, False, False, None),
    (38, "Worth it?", '=IF(ISNUMBER(B37),IF(B37<=B23,"Yes: under your max cost per door","No: improve the list, offer, or follow-up"),"-")', None, False, False, None),
]
for r, label, val, fmt, is_input, key, note in rows:
    if val is None and label:
        section(de, r, label)
        continue
    de.cell(row=r, column=1, value=label).font = label_font
    c = de.cell(row=r, column=2, value=val)
    c.font = input_font if is_input else calc_font
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    if is_input:
        c.fill = key_fill if key else input_fill
    if note:
        n = de.cell(row=r, column=3, value=note)
        n.font = note_font
        n.alignment = Alignment(wrap_text=True, vertical="top")
    for col in (1, 2, 3):
        de.cell(row=r, column=col).border = box
de["B20"].font = Font(name=F, size=10, bold=True)
de["B23"].font = Font(name=F, size=10, bold=True)
de["B26"].font = Font(name=F, size=10, bold=True)
de["B38"].alignment = Alignment(horizontal="left")
de.freeze_panes = "A4"

# ── Weekly Scoreboard ────────────────────────────────────────────────────────
ws = wb.create_sheet("Weekly Scoreboard")
ws["A1"], ws["A1"].font = "Weekly Scoreboard", title_font
ws["A2"], ws["A2"].font = "Fill one row every Monday. Blue cells are yours; black columns calculate.", sub_font
cols = ["Week of", "Website leads", "Rental analyses sent", "Consults booked", "Agreements signed", "Doors added",
        "Total doors", "Google reviews (total)", "Median reply time (hrs)", "Marketing spend ($)", "Spend to date ($)", "Cost per door to date ($)"]
header_row(ws, 4, cols, [14, 12, 13, 12, 13, 11, 11, 13, 14, 14, 14, 15])

# Example row (row 5): shows the format, excluded from totals.
example = [None, 6, 5, 3, 1, 1, "=F5", 4, 2, 250, "=J5", '=IF(G5>0,K5/G5,"-")']
ws.cell(row=5, column=1, value="Example").font = example_font
for col, v in enumerate(example[1:], 2):
    c = ws.cell(row=5, column=col, value=v)
    c.font = example_font
    c.number_format = USD if col in (10, 11, 12) else NUM
    c.alignment = Alignment(horizontal="center")

first, weeks = 6, 13
start = dt.date(2026, 9, 14)
for i in range(weeks):
    r = first + i
    a = ws.cell(row=r, column=1, value=start if i == 0 else f"=A{r-1}+7")
    a.number_format = "mmm d, yyyy"
    a.font = input_font if i == 0 else calc_font
    for col in (2, 3, 4, 5, 6, 8, 9, 10):
        c = ws.cell(row=r, column=col)
        c.font, c.fill = input_font, input_fill
        c.number_format = USD if col == 10 else NUM
        c.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=7, value=f"=F{r}" if i == 0 else f"=G{r-1}+F{r}")
    ws.cell(row=r, column=11, value=f"=J{r}" if i == 0 else f"=K{r-1}+J{r}")
    ws.cell(row=r, column=12, value=f'=IF(G{r}>0,K{r}/G{r},"-")')
    for col, fmt in ((7, NUM), (11, USD), (12, USD)):
        c = ws.cell(row=r, column=col)
        c.font, c.number_format = calc_font, fmt
        c.alignment = Alignment(horizontal="center")
    for col in range(1, 13):
        ws.cell(row=r, column=col).border = box

last = first + weeks - 1
tot = last + 1
ws.cell(row=tot, column=1, value="13-week total").font = Font(name=F, size=10, bold=True)
for col, formula, fmt in [
    (2, f"=SUM(B{first}:B{last})", NUM), (3, f"=SUM(C{first}:C{last})", NUM), (4, f"=SUM(D{first}:D{last})", NUM),
    (5, f"=SUM(E{first}:E{last})", NUM), (6, f"=SUM(F{first}:F{last})", NUM), (7, f"=G{last}", NUM),
    (8, f"=MAX(H{first}:H{last})", NUM), (9, f'=IF(COUNT(I{first}:I{last})>0,MEDIAN(I{first}:I{last}),"-")', NUM1),
    (10, f"=SUM(J{first}:J{last})", USD), (11, f"=K{last}", USD), (12, f'=IF(G{tot}>0,K{tot}/G{tot},"-")', USD),
]:
    c = ws.cell(row=tot, column=col, value=formula)
    c.font, c.number_format = Font(name=F, size=10, bold=True), fmt
    c.alignment = Alignment(horizontal="center")
    c.fill = section_fill
ws.cell(row=tot, column=1).fill = section_fill
ws.cell(row=tot + 2, column=1, value="Targets for the first 90 days: reply to every lead within 1 hour on business days; 10+ Google reviews; cost per door under the max on Door Economics.").font = note_font
ws.freeze_panes = "B5"

# ── Channels ────────────────────────────────────────────────────────────────
ch = wb.create_sheet("Channels")
ch["A1"], ch["A1"].font = "Channels", title_font
ch["A2"], ch["A2"].font = "Keep a running total per channel. Status compares cost per door with the max on Door Economics.", sub_font
header_row(ch, 4, ["Channel", "Spend ($)", "Leads", "Consults", "Signed doors", "Cost per lead ($)", "Cost per door ($)", "Lead → door rate", "Status", "Lead email tag"],
           [30, 12, 10, 11, 12, 14, 14, 13, 22, 36])
channels = [
    ("Google Business Profile", "google_business_profile / organic"),
    ("Postcards (absentee owners)", "postcard / direct_mail"),
    ("Agent referrals", "agent_onepager / print"),
    ("Switch outreach", "outreach / offline"),
    ("Social media (organic)", "facebook.com, instagram.com, linkedin.com"),
    ("Google Ads", "google ads"),
    ("Meta Ads", "meta"),
    ("Website / SEO", "google.com, bing.com"),
    ("Word of mouth / direct", "direct"),
]
for i, (name, tag) in enumerate(channels):
    r = 5 + i
    ch.cell(row=r, column=1, value=name).font = label_font
    for col in (2, 3, 4, 5):
        c = ch.cell(row=r, column=col)
        c.font, c.fill = input_font, input_fill
        c.number_format = USD if col == 2 else NUM
        c.alignment = Alignment(horizontal="center")
    ch.cell(row=r, column=6, value=f'=IF(C{r}>0,B{r}/C{r},"-")').number_format = USD
    ch.cell(row=r, column=7, value=f'=IF(E{r}>0,B{r}/E{r},"-")').number_format = USD
    ch.cell(row=r, column=8, value=f'=IF(C{r}>0,E{r}/C{r},"-")').number_format = PCT
    ch.cell(row=r, column=9, value=f"=IF(E{r}=0,IF(B{r}>0,\"No doors yet\",\"-\"),IF(B{r}/E{r}<='Door Economics'!$B$23,\"Scale up\",\"Fix or cut\"))")
    for col in (6, 7, 8, 9):
        c = ch.cell(row=r, column=col)
        c.font = calc_font
        c.alignment = Alignment(horizontal="center")
    ch.cell(row=r, column=10, value=tag).font = note_font
    for col in range(1, 11):
        ch.cell(row=r, column=col).border = box
lr = 5 + len(channels) - 1
tr = lr + 1
ch.cell(row=tr, column=1, value="Total").font = Font(name=F, size=10, bold=True)
for col, formula, fmt in [
    (2, f"=SUM(B5:B{lr})", USD), (3, f"=SUM(C5:C{lr})", NUM), (4, f"=SUM(D5:D{lr})", NUM), (5, f"=SUM(E5:E{lr})", NUM),
    (6, f'=IF(C{tr}>0,B{tr}/C{tr},"-")', USD), (7, f'=IF(E{tr}>0,B{tr}/E{tr},"-")', USD), (8, f'=IF(C{tr}>0,E{tr}/C{tr},"-")', PCT),
]:
    c = ch.cell(row=tr, column=col, value=formula)
    c.font, c.number_format = Font(name=F, size=10, bold=True), fmt
    c.alignment = Alignment(horizontal="center")
for col in range(1, 11):
    ch.cell(row=tr, column=col).fill = section_fill
ch["I4"].comment = Comment("Scale up = cost per door is at or under 'Max you can spend to win one door' (Door Economics!B23).", "TrueNorth")
ch.freeze_panes = "B5"

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

wb.save(OUT)
print(OUT)
